import sqlite3
import os
import csv
import io
from datetime import datetime, timedelta
from functools import wraps
from flask import Flask, render_template, request, redirect, url_for, flash, g, Response, session

app = Flask(__name__)
app.secret_key = "change-cette-cle-avant-la-demo"
DB_PATH = os.path.join(os.path.dirname(__file__), "catalogue.db")
DUREE_EMPRUNT_JOURS = 14
ADMIN_MOT_DE_PASSE = "biblio2026"


def get_db():
    if "db" not in g:
        g.db = sqlite3.connect(DB_PATH)
        g.db.row_factory = sqlite3.Row
        g.db.execute("PRAGMA foreign_keys = ON")
    return g.db


@app.teardown_appcontext
def close_db(exception=None):
    db = g.pop("db", None)
    if db is not None:
        db.close()


def admin_requis(vue):
    @wraps(vue)
    def wrapper(*args, **kwargs):
        if not session.get("est_admin"):
            flash("Connexion administrateur requise.", "error")
            return redirect(url_for("admin_login", next=request.path))
        return vue(*args, **kwargs)
    return wrapper


def _colonnes_existantes(db, table):
    return {row["name"] for row in db.execute(f"PRAGMA table_info({table})").fetchall()}


def init_db():
    db = sqlite3.connect(DB_PATH)
    db.row_factory = sqlite3.Row
    db.executescript(
        """
        CREATE TABLE IF NOT EXISTS livres (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            titre TEXT NOT NULL,
            auteur TEXT NOT NULL,
            categorie TEXT,
            disponible INTEGER NOT NULL DEFAULT 1
        );

        CREATE TABLE IF NOT EXISTS emprunts (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            livre_id INTEGER NOT NULL,
            emprunteur TEXT NOT NULL,
            statut TEXT NOT NULL DEFAULT 'en_attente',
            date_demande TEXT NOT NULL,
            date_emprunt TEXT,
            date_retour_prevue TEXT,
            date_retour_effective TEXT,
            FOREIGN KEY (livre_id) REFERENCES livres(id)
        );

        CREATE TABLE IF NOT EXISTS liste_attente (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            livre_id INTEGER NOT NULL,
            demandeur TEXT NOT NULL,
            date_demande TEXT NOT NULL,
            notifie INTEGER NOT NULL DEFAULT 0,
            FOREIGN KEY (livre_id) REFERENCES livres(id)
        );
        """
    )

    colonnes_emprunts = _colonnes_existantes(db, "emprunts")
    migrations = {
        "statut": "ALTER TABLE emprunts ADD COLUMN statut TEXT NOT NULL DEFAULT 'valide'",
        "date_demande": "ALTER TABLE emprunts ADD COLUMN date_demande TEXT",
        "telephone": "ALTER TABLE emprunts ADD COLUMN telephone TEXT",
        "email": "ALTER TABLE emprunts ADD COLUMN email TEXT",
    }
    for colonne, sql in migrations.items():
        if colonne not in colonnes_emprunts:
            db.execute(sql)
    if "date_demande" not in colonnes_emprunts:
        db.execute("UPDATE emprunts SET date_demande = date_emprunt WHERE date_demande IS NULL")
    db.commit()

    cur = db.execute("SELECT COUNT(*) FROM livres")
    if cur.fetchone()[0] == 0:
        db.executemany(
            "INSERT INTO livres (titre, auteur, categorie, disponible) VALUES (?, ?, ?, 1)",
            [
                ("Introduction a l'Algorithmique", "Cormen", "Informatique"),
                ("Reseaux Informatiques", "Tanenbaum", "Informatique"),
                ("Physique Generale", "Halliday", "Sciences"),
                ("Analyse Mathematique", "Rudin", "Mathematiques"),
                ("Bases de Donnees", "Elmasri", "Informatique"),
            ],
        )
        db.commit()
    db.close()


@app.route("/")
def index():
    db = get_db()
    q = request.args.get("q", "").strip()
    if q:
        livres = db.execute(
            "SELECT * FROM livres WHERE titre LIKE ? OR auteur LIKE ? ORDER BY titre",
            (f"%{q}%", f"%{q}%"),
        ).fetchall()
    else:
        livres = db.execute("SELECT * FROM livres ORDER BY titre").fetchall()

    livres_avec_infos = []
    for livre in livres:
        nb_attente = db.execute(
            "SELECT COUNT(*) FROM liste_attente WHERE livre_id = ?", (livre["id"],)
        ).fetchone()[0]
        demande_en_attente = db.execute(
            "SELECT id FROM emprunts WHERE livre_id = ? AND statut = 'en_attente'", (livre["id"],)
        ).fetchone()
        livres_avec_infos.append({
            **dict(livre),
            "nb_attente": nb_attente,
            "demande_en_attente": demande_en_attente is not None,
        })

    return render_template("index.html", livres=livres_avec_infos, q=q)


@app.route("/emprunter/<int:livre_id>", methods=["GET", "POST"])
def emprunter(livre_id):
    db = get_db()
    livre = db.execute("SELECT * FROM livres WHERE id = ?", (livre_id,)).fetchone()
    if livre is None:
        flash("Livre introuvable.", "error")
        return redirect(url_for("index"))
    if not livre["disponible"]:
        flash("Ce livre n'est plus disponible. Inscrivez-vous sur la liste d'attente depuis le catalogue.", "error")
        return redirect(url_for("index"))

    if request.method == "POST":
        emprunteur = request.form.get("emprunteur", "").strip()
        telephone = request.form.get("telephone", "").strip()
        email = request.form.get("email", "").strip()

        if not emprunteur:
            flash("Merci d'indiquer votre nom.", "error")
            return render_template("emprunter.html", livre=livre)
        if not telephone:
            flash("Le numero de telephone est obligatoire pour vous recontacter en cas de retard.", "error")
            return render_template("emprunter.html", livre=livre)

        db.execute(
            "INSERT INTO emprunts (livre_id, emprunteur, telephone, email, statut, date_demande) VALUES (?, ?, ?, ?, 'en_attente', ?)",
            (livre_id, emprunteur, telephone, email, datetime.now().isoformat()),
        )
        db.commit()
        flash("Demande envoyee ! Elle sera confirmee des qu'un responsable de la bibliotheque l'aura validee.", "success")
        return redirect(url_for("index"))

    return render_template("emprunter.html", livre=livre)


@app.route("/liste-attente/<int:livre_id>", methods=["POST"])
def rejoindre_liste_attente(livre_id):
    db = get_db()
    livre = db.execute("SELECT * FROM livres WHERE id = ?", (livre_id,)).fetchone()
    if livre is None:
        flash("Livre introuvable.", "error")
        return redirect(url_for("index"))

    demandeur = request.form.get("demandeur", "").strip()
    if not demandeur:
        flash("Merci d'indiquer votre nom pour rejoindre la liste d'attente.", "error")
        return redirect(url_for("index"))

    deja_inscrit = db.execute(
        "SELECT id FROM liste_attente WHERE livre_id = ? AND demandeur = ?",
        (livre_id, demandeur),
    ).fetchone()
    if deja_inscrit:
        flash("Vous etes deja sur la liste d'attente pour ce livre.", "error")
        return redirect(url_for("index"))

    db.execute(
        "INSERT INTO liste_attente (livre_id, demandeur, date_demande) VALUES (?, ?, ?)",
        (livre_id, demandeur, datetime.now().isoformat()),
    )
    db.commit()
    position = db.execute(
        "SELECT COUNT(*) FROM liste_attente WHERE livre_id = ?", (livre_id,)
    ).fetchone()[0]
    flash(f"Inscrit sur la liste d'attente pour '{livre['titre']}' (position {position}).", "success")
    return redirect(url_for("index"))


@app.route("/admin/login", methods=["GET", "POST"])
def admin_login():
    if request.method == "POST":
        mot_de_passe = request.form.get("mot_de_passe", "")
        if mot_de_passe == ADMIN_MOT_DE_PASSE:
            session["est_admin"] = True
            flash("Connecte en tant qu'administrateur.", "success")
            return redirect(request.args.get("next") or url_for("admin_demandes"))
        flash("Mot de passe incorrect.", "error")
    return render_template("admin_login.html")


@app.route("/admin/logout")
def admin_logout():
    session.pop("est_admin", None)
    flash("Deconnecte.", "success")
    return redirect(url_for("index"))


@app.route("/admin/demandes")
@admin_requis
def admin_demandes():
    db = get_db()
    demandes = db.execute(
        """
        SELECT emprunts.id, livres.titre, livres.auteur, emprunts.emprunteur,
               emprunts.telephone, emprunts.email, emprunts.date_demande
        FROM emprunts
        JOIN livres ON livres.id = emprunts.livre_id
        WHERE emprunts.statut = 'en_attente'
        ORDER BY emprunts.date_demande
        """
    ).fetchall()
    return render_template("admin_demandes.html", demandes=demandes)


@app.route("/admin/demandes/<int:emprunt_id>/approuver", methods=["POST"])
@admin_requis
def approuver_demande(emprunt_id):
    db = get_db()
    emprunt = db.execute("SELECT * FROM emprunts WHERE id = ?", (emprunt_id,)).fetchone()
    if emprunt is None or emprunt["statut"] != "en_attente":
        flash("Demande introuvable ou deja traitee.", "error")
        return redirect(url_for("admin_demandes"))

    livre = db.execute("SELECT * FROM livres WHERE id = ?", (emprunt["livre_id"],)).fetchone()
    if not livre["disponible"]:
        flash("Ce livre n'est plus disponible, impossible d'approuver.", "error")
        return redirect(url_for("admin_demandes"))

    maintenant = datetime.now()
    retour_prevu = maintenant + timedelta(days=DUREE_EMPRUNT_JOURS)
    db.execute(
        "UPDATE emprunts SET statut = 'valide', date_emprunt = ?, date_retour_prevue = ? WHERE id = ?",
        (maintenant.isoformat(), retour_prevu.isoformat(), emprunt_id),
    )
    db.execute("UPDATE livres SET disponible = 0 WHERE id = ?", (emprunt["livre_id"],))
    db.commit()
    flash(f"Emprunt approuve pour {emprunt['emprunteur']}.", "success")
    return redirect(url_for("admin_demandes"))


@app.route("/admin/demandes/<int:emprunt_id>/refuser", methods=["POST"])
@admin_requis
def refuser_demande(emprunt_id):
    db = get_db()
    emprunt = db.execute("SELECT * FROM emprunts WHERE id = ?", (emprunt_id,)).fetchone()
    if emprunt is None or emprunt["statut"] != "en_attente":
        flash("Demande introuvable ou deja traitee.", "error")
        return redirect(url_for("admin_demandes"))

    db.execute("UPDATE emprunts SET statut = 'refuse' WHERE id = ?", (emprunt_id,))
    db.commit()
    flash(f"Demande de {emprunt['emprunteur']} refusee.", "success")
    return redirect(url_for("admin_demandes"))


@app.route("/admin/emprunts")
@admin_requis
def admin_emprunts():
    db = get_db()
    en_cours = db.execute(
        """
        SELECT emprunts.id, livres.titre, livres.auteur, emprunts.emprunteur,
               emprunts.telephone, emprunts.email, emprunts.date_emprunt, emprunts.date_retour_prevue
        FROM emprunts
        JOIN livres ON livres.id = emprunts.livre_id
        WHERE emprunts.statut = 'valide' AND emprunts.date_retour_effective IS NULL
        ORDER BY emprunts.date_retour_prevue
        """
    ).fetchall()

    attente = db.execute(
        """
        SELECT liste_attente.id, livres.titre, liste_attente.demandeur, liste_attente.date_demande
        FROM liste_attente
        JOIN livres ON livres.id = liste_attente.livre_id
        ORDER BY livres.titre, liste_attente.date_demande
        """
    ).fetchall()

    return render_template("admin_emprunts.html", emprunts=en_cours, attente=attente)


@app.route("/retourner/<int:emprunt_id>", methods=["POST"])
@admin_requis
def retourner(emprunt_id):
    db = get_db()
    emprunt = db.execute("SELECT * FROM emprunts WHERE id = ?", (emprunt_id,)).fetchone()
    if emprunt is None:
        flash("Emprunt introuvable.", "error")
        return redirect(url_for("admin_emprunts"))

    db.execute(
        "UPDATE emprunts SET date_retour_effective = ? WHERE id = ?",
        (datetime.now().isoformat(), emprunt_id),
    )
    db.execute("UPDATE livres SET disponible = 1 WHERE id = ?", (emprunt["livre_id"],))
    db.commit()

    prochain = db.execute(
        "SELECT * FROM liste_attente WHERE livre_id = ? ORDER BY date_demande LIMIT 1",
        (emprunt["livre_id"],),
    ).fetchone()
    if prochain:
        flash(
            f"Livre rendu ! {prochain['demandeur']} est en tete de la liste d'attente pour ce livre.",
            "success",
        )
    else:
        flash("Livre rendu, merci !", "success")
    return redirect(url_for("admin_emprunts"))


@app.route("/liste-attente/retirer/<int:attente_id>", methods=["POST"])
@admin_requis
def retirer_liste_attente(attente_id):
    db = get_db()
    db.execute("DELETE FROM liste_attente WHERE id = ?", (attente_id,))
    db.commit()
    flash("Retire de la liste d'attente.", "success")
    return redirect(url_for("admin_emprunts"))


def _recuperer_lignes_export(db):
    return db.execute(
        """
        SELECT livres.titre, livres.auteur, emprunts.emprunteur, emprunts.telephone, emprunts.email, emprunts.statut,
               emprunts.date_demande, emprunts.date_emprunt, emprunts.date_retour_prevue, emprunts.date_retour_effective
        FROM emprunts
        JOIN livres ON livres.id = emprunts.livre_id
        ORDER BY emprunts.date_demande DESC
        """
    ).fetchall()


@app.route("/export/emprunts.csv")
@admin_requis
def export_emprunts_csv():
    db = get_db()
    lignes = _recuperer_lignes_export(db)

    buffer = io.StringIO()
    writer = csv.writer(buffer)
    writer.writerow(["Titre", "Auteur", "Emprunteur", "Telephone", "Email", "Statut", "Date demande", "Date emprunt", "Date retour prevue", "Date retour effective"])
    for l in lignes:
        writer.writerow([
            l["titre"], l["auteur"], l["emprunteur"], l["telephone"] or "", l["email"] or "", l["statut"],
            l["date_demande"][:16].replace("T", " ") if l["date_demande"] else "",
            l["date_emprunt"][:16].replace("T", " ") if l["date_emprunt"] else "",
            l["date_retour_prevue"][:10] if l["date_retour_prevue"] else "",
            l["date_retour_effective"][:16].replace("T", " ") if l["date_retour_effective"] else "",
        ])

    return Response(
        buffer.getvalue(),
        mimetype="text/csv",
        headers={"Content-Disposition": "attachment; filename=emprunts_catalogue_plus.csv"},
    )


@app.route("/export/emprunts.xlsx")
@admin_requis
def export_emprunts_xlsx():
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter

    db = get_db()
    lignes = _recuperer_lignes_export(db)

    wb = Workbook()
    ws = wb.active
    ws.title = "Emprunts"

    entetes = ["Titre", "Auteur", "Emprunteur", "Telephone", "Email", "Statut", "Date demande", "Date emprunt", "Date retour prevue", "Date retour effective"]
    ws.append(entetes)
    for col_idx in range(1, len(entetes) + 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill(start_color="4F46E5", end_color="4F46E5", fill_type="solid")
        cell.alignment = Alignment(horizontal="center")

    statut_labels = {"en_attente": "En attente", "valide": "Valide", "refuse": "Refuse"}
    for l in lignes:
        ws.append([
            l["titre"], l["auteur"], l["emprunteur"], l["telephone"] or "", l["email"] or "",
            statut_labels.get(l["statut"], l["statut"]),
            l["date_demande"][:16].replace("T", " ") if l["date_demande"] else "",
            l["date_emprunt"][:16].replace("T", " ") if l["date_emprunt"] else "",
            l["date_retour_prevue"][:10] if l["date_retour_prevue"] else "",
            l["date_retour_effective"][:16].replace("T", " ") if l["date_retour_effective"] else "",
        ])

    largeurs = [32, 20, 20, 16, 24, 12, 18, 18, 18, 20]
    for i, largeur in enumerate(largeurs, start=1):
        ws.column_dimensions[get_column_letter(i)].width = largeur

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)

    return Response(
        buffer.getvalue(),
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": "attachment; filename=emprunts_catalogue_plus.xlsx"},
    )


if __name__ == "__main__":
    init_db()
    print("Serveur local demarre.")
    print("Sur ce PC : http://127.0.0.1:5000")
    print("Depuis un telephone sur le meme wifi : http://<IP-DE-CE-PC>:5000")
    print("Espace admin : /admin/login (mot de passe par defaut : biblio2026)")
    app.run(host="0.0.0.0", port=5000, debug=True)
